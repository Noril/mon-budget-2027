"""Eurobaromètre standard, volume A (un classeur XLSX par vague) -> confiance dans les institutions, format long.

Chaque vague publie un onglet par item de la question « How much trust do you have in certain institutions? »
(QA6.x ; le numéro de l'item change d'une vague à l'autre). Dans chaque onglet : ligne 1 « Eurobarometer - 105.2 »,
ligne 2 « Terrain/Fieldwork : 12/3 - 5/4/2026 », ligne 4 l'item en français (colonne B) et en anglais, puis
une ligne d'en-tête (UE27, BE, BG…) et les lignes « Total » (effectif pondéré), « Tend to trust »,
« Tend not to trust », « Don't know » (parts entre 0 et 1, arrondies au centième par la Commission ; « - » = 0).

Sortie : une ligne par vague, item et pays. `institution` normalise l'item anglais : gouvernement
(« Government »), parlement (« PARLIAMENT » national), justice (« legal system »), union_europeenne,
sinon vide. Pays : code de l'en-tête, UE27 -> EU27_2020. `fin_terrain` : dernier jour du terrain.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import duckdb
import openpyxl

QUESTION = "How much trust do you have in certain institutions"
TERRAIN = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})\s*$")


def institution(item_en: str) -> str | None:
    bas = item_en.lower()
    if "legal system" in bas:
        return "justice"
    if "government" in bas:
        return "gouvernement"
    if "parliament" in bas and "european" not in bas:
        return "parlement"
    if "european union" in bas:
        return "union_europeenne"
    return None


def part(valeur) -> float | None:
    if valeur in (None, ""):
        return None
    return 0.0 if valeur == "-" else float(valeur)


def lire_classeur(f: Path) -> list[tuple]:
    lignes = []
    classeur = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for feuille in classeur.sheetnames:
        rangs = [list(r) for r in classeur[feuille].iter_rows(max_row=30, values_only=True)]
        if len(rangs) < 12:
            continue
        colonne_en = next((i for i, c in enumerate(rangs[2]) if isinstance(c, str) and QUESTION in c), None)
        if colonne_en is None:
            continue
        vague = str(next(c for c in rangs[0] if c)).replace("Eurobarometer -", "").strip()
        terrain = " ".join(str(c) for c in rangs[1] if isinstance(c, str) and "Fieldwork" in c)
        m = TERRAIN.search(terrain)
        if not m:
            raise ValueError(f"{f.name}/{feuille} : terrain illisible {terrain!r}")
        fin = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        item_fr, item_en = str(rangs[3][1] or ""), str(rangs[3][colonne_en] or "")
        (entete,) = [r for r in rangs if isinstance(r[2], str) and r[2].startswith("UE27")]
        par_libelle = {r[1]: r for r in rangs if isinstance(r[1], str)}
        for j in range(2, len(entete)):
            if not entete[j]:
                continue
            pays = "EU27_2020" if str(entete[j]).startswith("UE27") else str(entete[j]).strip()
            lignes.append((
                vague, terrain.split(":", 1)[-1].strip(), fin, feuille, item_fr, item_en, institution(item_en), pays,
                part(par_libelle["Tend to trust"][j]), part(par_libelle["Tend not to trust"][j]),
                part(par_libelle["Don't know"][j]), par_libelle["Total"][j], f.name,
            ))
    return lignes


def normaliser(source: dict, fichiers: list[Path], sortie: Path) -> None:
    lignes = [ligne for f in fichiers for ligne in lire_classeur(f)]
    if not lignes:
        raise ValueError(f"{source['id']} : aucune question de confiance trouvée")
    con = duckdb.connect()
    con.execute(
        "CREATE TABLE t (vague VARCHAR, terrain VARCHAR, fin_terrain DATE, feuille VARCHAR, item_fr VARCHAR, "
        "item_en VARCHAR, institution VARCHAR, pays VARCHAR, confiance DOUBLE, pas_confiance DOUBLE, nsp DOUBLE, "
        "effectif DOUBLE, fichier VARCHAR)"
    )
    con.executemany("INSERT INTO t VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", lignes)
    con.execute(f"COPY t TO '{sortie}' (FORMAT parquet)")
    con.close()
